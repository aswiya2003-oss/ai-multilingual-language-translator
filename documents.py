from io import BytesIO
from pathlib import Path

def extract_text(filename, content):
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        for encoding in ("utf-8", "utf-8-sig", "cp1252"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                pass
        raise ValueError("Could not decode text file. Save it as UTF-8 and retry.")
    if suffix == ".docx":
        from docx import Document
        document = Document(BytesIO(content))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n\n".join(parts)
    if suffix == ".pdf":
        import fitz
        doc = fitz.open(stream=content, filetype="pdf")
        return "\n\n".join(page.get_text("text") for page in doc)
    raise ValueError("Unsupported document type.")

def make_docx(text):
    from docx import Document
    doc = Document()
    for paragraph in text.split("\n"):
        doc.add_paragraph(paragraph)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
